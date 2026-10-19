"""Copy the economy's levers from 'Stat Street Economy.xlsx' into config/economy.json, which the game reads.

Usage (from sports-exchange/):  python3 tools/export_economy.py
Change a blue cell in Excel, save, run this, and the game follows. Only input cells are read, never formulas, so the
workbook doesn't need recalculating first. Cells are found by their labels, so moving rows around is fine.
"""
import json
import os
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BOOK = os.path.join(ROOT, "Stat Street Economy.xlsx")
OUT = os.path.join(ROOT, "config", "economy.json")


def by_label(ws, label, col=2):
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == label:
            v = ws.cell(row=row[0].row, column=col).value
            if not isinstance(v, (int, float)):
                sys.exit(f"'{ws.title}' row '{label}': expected a number, found {v!r}")
            return v
    sys.exit(f"Couldn't find '{label}' on the {ws.title} sheet.")


def export(book=BOOK):
    wb = load_workbook(book)
    st, sc, mm, live = wb["Settings"], wb["Scoring"], wb["Market Maker"], wb["Live Rules"]
    out = {
        "starting_cash": by_label(st, "Starting cash per player"),
        "target_price": by_label(st, "Target average starting price"),
        "weekly_yield": by_label(st, "Weekly dividend yield (in season)"),
        "sports": {}, "groups": {},
        "market_maker": {
            "spread": by_label(mm, "Buy-sell gap (spread)"),
            "lean_per_step": by_label(mm, "Price lean per step of league holdings"),
            "dollars_per_step": by_label(mm, "Dollars of fair value per step"),
        },
        "live": {
            "blend_games": by_label(live, "Blend weight (games)"),
            "recheck_weeks": by_label(live, "Re-check weight (weeks)"),
            "recheck_max_move": by_label(live, "Re-check: most it can move in a week"),
            "recheck_floor": by_label(live, "Re-check: floor"),
            "recheck_ceiling": by_label(live, "Re-check: ceiling"),
            "playoff_boost": by_label(live, "Playoff dividend boost"),
        },
    }
    # Per sport: the row whose first cell is the sport's key, under the "Per sport" header.
    for row in st.iter_rows(min_row=1, max_row=st.max_row):
        key = row[0].value
        if key in ("NFL", "NBA", "NHL") and isinstance(row[1].value, (int, float)) and key not in out["sports"]:
            out["sports"][key] = {"weeks": row[1].value, "games": row[2].value, "min_games": row[3].value,
                                  "listed": row[4].value, "scoring": {}}
        # Position groups: name, sport, positions, lean, repeat.
        if isinstance(row[1].value, str) and row[1].value in ("NFL", "NBA", "NHL") and isinstance(row[3].value, float) \
                and isinstance(row[4].value, float):
            out["groups"][key] = {"sport": row[1].value, "positions": row[2].value, "lean": row[3].value,
                                  "repeat": row[4].value}
    for row in sc.iter_rows(min_row=5, max_col=4):
        sport, _, pts, col = (c.value for c in row)
        if sport in out["sports"] and col:
            out["sports"][sport]["scoring"][col] = pts
    if set(out["sports"]) != {"NFL", "NBA", "NHL"} or len(out["groups"]) != 6:
        sys.exit(f"Workbook layout not recognised: sports {sorted(out['sports'])}, groups {sorted(out['groups'])}")
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w") as f:
        json.dump(out, f, indent=2)
        f.write("\n")
    return out


if __name__ == "__main__":
    export()
    print(f"wrote {os.path.relpath(OUT, ROOT)}")
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import build_rules  # the rules page quotes these numbers, so rebuild it too
    build_rules.main()
