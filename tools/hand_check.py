"""The weekly hand-check workbook: python3 tools/hand_check.py --week 6   (or no --week: the latest week paid)

Writes checks/nfl-2026-week-NN.xlsx. It holds ESPN's raw stats for every listed athlete in every game of the week,
the scoring weights and dividend rate as blue input cells, your trades, and formulas that score each stat line,
work out the dividend per share and your payout, and compare them with what the game paid. The formulas don't read
the game's answers; they only compare against them, so a "match" means two separate calculations agree.

Called by tools/nightly.py after each day it processes. Options: --save PATH, --out FOLDER, --no-recalc.

For a player in an online league: python3 tools/hand_check.py --online --target test --league CODE --player "Name"
It reads that player's trades and dividends from the database, so player-to-player trades count, and writes
checks/online-CODE-Name/. The stat lines come from the database too: they are what the robot paid on.
"""
import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as col

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from engine import api, ledger as L  # noqa: E402

DEFAULT_SAVE = os.path.join(ROOT, "save", "game.sqlite")
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", start_color="1F3864")
BAD_FILL = PatternFill("solid", start_color="FFC7CE")
MONEY = '$#,##0.00;($#,##0.00);"-"'
EPOCH = datetime(1899, 12, 30, tzinfo=timezone.utc)


def serial(utc_text):
    """An Excel date-time number for a UTC time, so formulas can compare times."""
    return round((L.parse_utc(utc_text) - EPOCH).total_seconds() / 86400, 8)


def latest_week(conn):
    row = conn.execute("SELECT MAX(week) FROM games WHERE processed_at IS NOT NULL").fetchone()
    return row[0]


def write_week(conn, week, out_dir):
    """Build the workbook for one week; returns its path, or None if no game of that week has been paid."""
    sport, season = L.meta(conn, "sport"), L.meta(conn, "season")
    econ = L.economy(conn)
    weights = econ["sports"][sport]["scoring"]
    stats = list(weights)
    rate = L.meta(conn, "rate")
    boost = econ["live"]["playoff_boost"]
    lines = conn.execute("""
        SELECT g.id AS gid, g.day, g.home, g.away, g.playoff, g.kickoff, a.ticker, a.name, s.team, s.stats, s.points,
               p.per_share, p.record_time, COALESCE(d.shares, 0) AS paid_shares, COALESCE(d.cents, 0) AS paid_cents
        FROM games g JOIN stat_lines s ON s.game_id = g.id JOIN athletes a ON a.id = s.athlete_id
        JOIN payouts p ON p.game_id = g.id AND p.athlete_id = s.athlete_id
        LEFT JOIN dividends d ON d.game_id = g.id AND d.athlete_id = s.athlete_id AND d.player = ?
        WHERE g.week = ? AND g.processed_at IS NOT NULL ORDER BY g.kickoff, g.id, a.ticker""", (api.PLAYER, week)).fetchall()
    if not lines:
        return None
    trades = conn.execute("""SELECT t.at, a.ticker, t.shares, t.cents FROM trades t JOIN athletes a ON a.id = t.athlete_id
                             WHERE t.player = ? ORDER BY t.at, t.id""", (api.PLAYER,)).fetchall()

    wb = Workbook()
    ck, sc, tr = wb.active, wb.create_sheet("Scoring"), wb.create_sheet("Trades")
    ck.title = "Check"

    # Scoring: the weights and the rate, as blue inputs.
    sc["A1"], sc["A2"] = "Scoring weights (points per stat)", "Change a blue cell and every check recalculates."
    sc["A1"].font, sc["A2"].font = BOLD, BLACK
    for i, k in enumerate(stats):
        r = 4 + i
        sc.cell(r, 1, k).font = BLACK
        c = sc.cell(r, 2, weights[k])
        c.font = BLUE
    r_rate, r_boost = 5 + len(stats), 6 + len(stats)
    sc.cell(r_rate, 1, "Dividend rate (dollars per fantasy point, per share)").font = BLACK
    sc.cell(r_rate, 2, rate).font = BLUE
    sc.cell(r_boost, 1, "Playoff boost").font = BLACK
    sc.cell(r_boost, 2, boost).font = BLUE
    sc.cell(r_rate, 2).number_format = "0.000000"
    sc.column_dimensions["A"].width = 52
    wrow = {k: 4 + i for i, k in enumerate(stats)}

    # Trades: every trade you made, so holdings at each record time are added up here, not copied from the game.
    for j, h in enumerate(["Time (UTC)", "Time", "Ticker", "Shares (negative = sold)", "Cents"], 1):
        c = tr.cell(1, j, h)
        c.font, c.fill = HEAD, HEAD_FILL
    for i, t in enumerate(trades, 2):
        tr.cell(i, 1, serial(t["at"])).number_format = "yyyy-mm-dd hh:mm:ss"
        tr.cell(i, 2, t["at"])
        tr.cell(i, 3, t["ticker"])
        tr.cell(i, 4, t["shares"])
        tr.cell(i, 5, t["cents"])
        for j in range(1, 6):
            tr.cell(i, j).font = BLUE
    last_t = max(len(trades) + 1, 2)
    for j, w in enumerate([20, 22, 10, 24, 12], 1):
        tr.column_dimensions[col(j)].width = w

    # Check: one row per athlete per game.
    heads = ["Game", "Day", "Ticker", "Athlete", "Playoff? (1 = yes)", "Record time (UTC)"] + stats + [
        "Points", "Per share", "Shares you held", "Your payout",
        "Game: points", "Game: per share", "Game: shares held", "Game: payout", "Result"]
    first = 8
    hrow = first - 1
    for j, h in enumerate(heads, 1):
        c = ck.cell(hrow, j, h)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ns = len(stats)
    c_play, c_rec, c_s0 = 5, 6, 7
    c_pts, c_ps, c_sh, c_pay = c_s0 + ns, c_s0 + ns + 1, c_s0 + ns + 2, c_s0 + ns + 3
    g_pts, g_ps, g_sh, g_pay, c_res = c_pay + 1, c_pay + 2, c_pay + 3, c_pay + 4, c_pay + 5
    last = first + len(lines) - 1
    for i, ln in enumerate(lines):
        r = first + i
        raw = json.loads(ln["stats"])
        vals = [f"{ln['away']} at {ln['home']}", ln["day"], ln["ticker"], ln["name"], ln["playoff"],
                serial(ln["record_time"])] + [raw.get(k, 0) for k in stats]
        for j, v in enumerate(vals, 1):
            c = ck.cell(r, j, v)
            c.font = BLUE if j >= 5 else BLACK
        ck.cell(r, c_rec).number_format = "yyyy-mm-dd hh:mm"
        pts = "+".join(f"{col(c_s0 + n)}{r}*Scoring!$B${wrow[k]}" for n, k in enumerate(stats))
        ck.cell(r, c_pts, f"=ROUND({pts},2)")
        ck.cell(r, c_ps, f"=MAX(0,{col(c_pts)}{r})*Scoring!$B${r_rate}*IF({col(c_play)}{r}=1,Scoring!$B${r_boost},1)")
        ck.cell(r, c_sh, f"=SUMIFS(Trades!$D$2:$D${last_t},Trades!$C$2:$C${last_t},C{r},"
                         f"Trades!$A$2:$A${last_t},\"<\"&{col(c_rec)}{r})")
        ck.cell(r, c_pay, f"=ROUND({col(c_sh)}{r}*{col(c_ps)}{r},2)")
        for j, v in ((g_pts, ln["points"]), (g_ps, ln["per_share"]), (g_sh, ln["paid_shares"]),
                     (g_pay, ln["paid_cents"] / 100)):
            ck.cell(r, j, v).font = BLUE
        ck.cell(r, c_res, f"=IF(AND(ABS({col(c_pts)}{r}-{col(g_pts)}{r})<0.005,"
                          f"ABS({col(c_ps)}{r}-{col(g_ps)}{r})<0.000001,{col(c_sh)}{r}={col(g_sh)}{r},"
                          f"ABS({col(c_pay)}{r}-{col(g_pay)}{r})<0.005),\"match\",\"MISMATCH\")")
        for j in (c_pay, g_pay):
            ck.cell(r, j).number_format = MONEY
        for j in (c_pts, c_ps, g_pts, g_ps, c_sh, c_pay, c_res):
            ck.cell(r, j).font = BLACK
        ck.cell(r, c_ps).number_format = ck.cell(r, g_ps).number_format = "0.0000"

    res = f"{col(c_res)}{first}:{col(c_res)}{last}"
    ck["A1"] = f"Hand check: {sport} {season} week {week}"
    ck["A1"].font = Font(name=FONT, bold=True, size=14)
    top = [("Lines checked", f"=COUNTA({res})", "0"),
           ("Mismatches", f'=COUNTIF({res},"MISMATCH")', "0"),
           ("Dividends the game paid you", f"=SUM({col(g_pay)}{first}:{col(g_pay)}{last})", MONEY),
           ("Dividends recalculated here", f"=SUM({col(c_pay)}{first}:{col(c_pay)}{last})", MONEY)]
    for i, (label, f, fmt) in enumerate(top):
        ck.cell(2 + i, 1, label).font = BOLD
        c = ck.cell(2 + i, 3, f)
        c.number_format, c.font = fmt, BOLD
    ck.cell(6, 1, "Blue = typed in or copied from the game or ESPN; black = formula. Weights and rate are on the Scoring sheet.").font = BLACK
    ck.freeze_panes = ck.cell(first, 5)
    ck.row_dimensions[hrow].height = 45
    for j, w in enumerate([16, 11, 8, 22, 10, 17] + [8] * ns + [9, 10, 10, 12, 10, 10, 10, 12, 11], 1):
        ck.column_dimensions[col(j)].width = w
    from openpyxl.formatting.rule import CellIsRule
    ck.conditional_formatting.add(res, CellIsRule(operator="equal", formula=['"MISMATCH"'], fill=BAD_FILL))

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"{sport.lower()}-{season}-week-{week:02d}.xlsx")
    wb.save(path)
    return path


def recalculate(path):
    """Have LibreOffice calculate every formula and store the results (so previews show numbers). Returns True if done."""
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if not os.path.exists(soffice):
        return False
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", f"-env:UserInstallation=file://{tmp}/profile", "--convert-to", "xlsx",
                        "--outdir", tmp, path], capture_output=True, timeout=180)
        done = os.path.join(tmp, os.path.basename(path))
        if not os.path.exists(done):
            return False
        shutil.copyfile(done, path)
    return True


def results(path):
    """(lines, mismatches) read back from a recalculated workbook."""
    from openpyxl import load_workbook
    ck = load_workbook(path, data_only=True)["Check"]
    return ck["C2"].value, ck["C3"].value


def online_save(project, season, market, member):
    """One league player's money in an in-memory save, in the Milestone 1 shape write_week reads."""
    from engine import robot
    conn, _ = robot.mirror(project, season)
    espn = lambda online_id: online_id.split(":", 1)[1]  # noqa: E731
    conn.execute("BEGIN")
    for p in project.select("payouts", game_id=f"like.{season}:*"):
        conn.execute("INSERT INTO payouts VALUES (?,?,?,?,?,?,?)",
                     (espn(p["game_id"]), espn(p["athlete_id"]), p["points"], p["boost"], p["per_share"],
                      L.utc_text(L.parse_utc(p["record_time"][:19])), L.utc_text(L.parse_utc(p["paid_at"][:19]))))
    for d in project.select("dividends", market_id=f"eq.{market}", member_id=f"eq.{member}"):
        conn.execute("INSERT INTO dividends VALUES (?,?,?,?,?)",
                     (espn(d["game_id"]), espn(d["athlete_id"]), api.PLAYER, d["shares"], d["cents"]))
    rows = project.select("trades", market_id=f"eq.{market}", order="at",
                          **{"or": f"(buyer_id.eq.{member},seller_id.eq.{member})"})
    for t in rows:
        sign = 1 if t["buyer_id"] == member else -1
        conn.execute("INSERT INTO trades (at, player, athlete_id, shares, cents, fair_value, market_held_before) "
                     "VALUES (?,?,?,?,?,0,0)", (L.utc_text(L.parse_utc(t["at"][:19])), api.PLAYER, espn(t["athlete_id"]),
                                                sign * t["shares"], sign * t["cents"]))
    conn.execute("COMMIT")
    return conn


def online_main(args):
    from engine.online import Project
    p = Project(args.target)
    league = p.select("markets", join_code=f"eq.{args.league}")
    if not league:
        sys.exit(f"No league with code {args.league} on {args.target}.")
    member = [m for m in p.select("members", market_id=f"eq.{league[0]['id']}")
              if m["display_name"].lower() == args.player.lower()]
    if not member:
        sys.exit(f"Nobody called {args.player} in that league.")
    conn = online_save(p, league[0]["seasons"][0], league[0]["id"], member[0]["id"])
    week = args.week or latest_week(conn)
    out = os.path.join(args.out, f"online-{args.league}-{member[0]['display_name'].replace(' ', '_')}")
    path = write_week(conn, week, out) if week else None
    if not path:
        sys.exit("No paid games for that week yet.")
    if not args.no_recalc and recalculate(path):
        n, bad = results(path)
        print(f"{path}: {n} lines, {bad} mismatches")
    else:
        print(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--online", action="store_true")
    ap.add_argument("--target", choices=["test", "live"], default="test")
    ap.add_argument("--league", help="the league's join code, with --online")
    ap.add_argument("--player", help="a display name in that league, with --online")
    ap.add_argument("--week", type=int)
    ap.add_argument("--save", default=DEFAULT_SAVE)
    ap.add_argument("--out", default=os.path.join(ROOT, "checks"))
    ap.add_argument("--no-recalc", action="store_true")
    args = ap.parse_args()
    if args.online:
        return online_main(args)
    conn = L.connect(args.save)
    week = args.week or latest_week(conn)
    path = write_week(conn, week, args.out) if week else None
    if not path:
        sys.exit("No paid games for that week yet.")
    if not args.no_recalc and recalculate(path):
        n, bad = results(path)
        print(f"{path}: {n} lines, {bad} mismatches")
    else:
        print(path)


if __name__ == "__main__":
    main()
