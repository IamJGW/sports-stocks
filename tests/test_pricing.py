import os
import unittest

from helpers import ECON, ROOT
from engine import pricing


class WorkbookAgreement(unittest.TestCase):
    """The code's starting prices must equal the workbook's (on the workbook's own seasons)."""

    def test_matches_recalculated_workbook(self):
        from openpyxl import load_workbook
        wb = load_workbook(os.path.join(ROOT, "Stat Street Economy.xlsx"), data_only=True, read_only=True)
        rate = wb["Settings"]["I12"].value
        if rate is None:
            self.skipTest("Workbook isn't recalculated (open and save it in Excel, or run the recalc script).")
        rows = list(wb["NFL"].iter_rows(min_row=2, values_only=True))
        hdr = rows[0]
        name, price, listed = hdr.index("Athlete"), hdr.index("Starting price"), hdr.index("Listed (1=yes)")
        book_prices = {r[name]: r[price] for r in rows[1:] if r[listed] == 1}
        book = pricing.starting_book(pricing.load_csv(os.path.join(ROOT, "data", "nfl_2024.csv")), ECON, "NFL")
        self.assertAlmostEqual(book["rate"], rate, places=12)
        ours = {a["name"]: a["price"] for a in book["athletes"]}
        self.assertEqual(set(ours), set(book_prices))
        for n, p in book_prices.items():
            self.assertAlmostEqual(ours[n], p, places=9, msg=n)


class InSeason(unittest.TestCase):
    def test_blend(self):
        k = ECON["live"]["blend_games"]
        self.assertEqual(pricing.expected_ptg(10, 1.0, 0, 0, ECON), 10)          # no games yet: the prior
        self.assertAlmostEqual(pricing.expected_ptg(10, 1.0, 20 * k, k, ECON), 15)  # k games at 20: halfway
        self.assertAlmostEqual(pricing.expected_ptg(10, 1.1, 0, 0, ECON), 11)    # multiplier scales the prior

    def test_recheck_speed_limit_and_bounds(self):
        lv = ECON["live"]
        new, ratio = pricing.recheck(1.0, 1, 150, 100, ECON)   # one wild week
        self.assertEqual(ratio, 1.5)
        self.assertAlmostEqual(new, 1 + lv["recheck_max_move"])
        new, _ = pricing.recheck(1.0, 10, 101, 100, ECON)      # a small, steady edge moves it a little
        self.assertAlmostEqual(new, (10 * 1.01 + lv["recheck_weeks"]) / (10 + lv["recheck_weeks"]))
        m = 1.0
        for week in range(1, 40):
            m, _ = pricing.recheck(m, week, 300, 100, ECON)
        self.assertEqual(m, lv["recheck_ceiling"])
        m = 1.0
        for week in range(1, 40):
            m, _ = pricing.recheck(m, week, 0, 100, ECON)
        self.assertEqual(m, lv["recheck_floor"])

    def test_average_starting_price_is_the_target(self):
        book = pricing.starting_book(pricing.load_csv(os.path.join(ROOT, "data", "nfl_2025.csv")), ECON, "NFL")
        self.assertEqual(len(book["athletes"]), 120)
        avg = sum(a["price"] for a in book["athletes"]) / 120
        self.assertAlmostEqual(avg, ECON["target_price"], places=6)


if __name__ == "__main__":
    unittest.main()
