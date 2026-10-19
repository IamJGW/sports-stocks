import os
import shutil
import tempfile
import unittest

from openpyxl import load_workbook

from helpers import FakeFeed, eastern, new_market
from engine import ledger as L, nightly as N
from tools import hand_check as H


@unittest.skipUnless(shutil.which("soffice") or os.path.exists("/Applications/LibreOffice.app"), "needs LibreOffice")
class HandCheck(unittest.TestCase):
    def setUp(self):
        self.c = new_market()
        rows = self.c.execute("SELECT * FROM athletes ORDER BY start_price DESC").fetchall()
        a = rows[0]
        b = next(r for r in rows if r["team"] != a["team"])
        feed = FakeFeed()
        feed.add_game("g1", 1, eastern(2026, 9, 13, 13), a["team"], b["team"], {
            a["id"]: (a["team"], {"rush_yd": 101, "rush_td": 2, "rec": 3, "rec_yd": 25, "fum_lost": 1}),
            b["id"]: (b["team"], {"int": 3, "fum_lost": 1}),
        })
        L.trade(self.c, "you", a["ticker"], "buy", 137, eastern(2026, 9, 10))
        L.trade(self.c, "you", a["ticker"], "sell", 37, eastern(2026, 9, 12, 20))
        L.trade(self.c, "you", a["ticker"], "buy", 500, eastern(2026, 9, 13, 0, 1))       # after record time: not paid
        L.trade(self.c, "you", b["ticker"], "buy", 61, eastern(2026, 9, 11))
        now = eastern(2026, 9, 14, 4)
        N.sync_schedule(self.c, feed, weeks=[1])
        for day in N.days_to_process(self.c, L.eastern_day(now)):
            N.process_day(self.c, day, N.fetch_day(self.c, day, feed), now)
        self.out = tempfile.mkdtemp()

    def test_replayed_week_has_no_mismatches(self):
        path = H.write_week(self.c, 1, self.out)
        self.assertTrue(H.recalculate(path))
        lines, bad = H.results(path)
        self.assertEqual((lines, bad), (2, 0))
        ck = load_workbook(path, data_only=True)["Check"]
        self.assertEqual(ck["C4"].value, ck["C5"].value)       # game paid = recalculated
        self.assertGreater(ck["C4"].value, 0)

    def test_workbook_is_live_so_a_wrong_weight_shows_up(self):
        path = H.write_week(self.c, 1, self.out)
        wb = load_workbook(path)
        wb["Scoring"]["B4"] = 0.05           # change the passing-yard weight
        wb["Scoring"]["B7"] = 7              # and rushing TDs
        wb.save(path)
        H.recalculate(path)
        self.assertGreater(H.results(path)[1], 0)

    def test_no_paid_games_no_workbook(self):
        self.assertIsNone(H.write_week(self.c, 2, self.out))


if __name__ == "__main__":
    unittest.main()
